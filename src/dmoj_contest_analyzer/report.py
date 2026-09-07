from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
HEADER_FONT = Font(bold=True)


def _safe(value):
    if isinstance(value, str) and value.lstrip()[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def _write_sheet(ws, headers, rows):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for row in rows:
        out = []
        for h in headers:
            v = row.get(h, "")
            out.append("" if v is None else _safe(v))
        ws.append(out)
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, start=1):
        max_len = max([len(str(h))] + [len(str(row.get(h, ""))) for row in rows]) if rows else len(h)  # noqa: E501
        ws.column_dimensions[get_column_letter(i)].width = min(max_len + 2, 45)


def write_excel_report(data, out_path: Path):
    main_rows, jplag_rows = data.main_rows, data.jplag_rows
    n_subs, n_users, n_problems = data.n_subs, data.n_users, data.n_problems
    wb = Workbook()

    ws_resumen = wb.active
    ws_resumen.title = "Resumen"
    ws_resumen.append(["Métrica", "Valor"])
    for cell in ws_resumen[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    n_alertas = sum(1 for r in main_rows if r["score_sospecha"] >= 2)
    n_jplag_altas = sum(1 for r in jplag_rows if r["similitud"] >= 70)
    resumen_data = [
        ("Total de submissions procesadas", n_subs),
        ("Usuarios distintos", n_users),
        ("Problemas distintos", n_problems),
        ("Casos con score_sospecha >= 2 (revisión prioritaria)", n_alertas),
        ("Pares JPlag con similitud >= 70%", n_jplag_altas),
    ]
    if data.llm_model:
        llm_rows = data.llm_rows or []
        resumen_data.extend([
            ("Modelo LLM usado", data.llm_model),
            ("Envíos juzgados por LLM", len(llm_rows)),
            ("Casos con llm_ai_score >= 70",
             sum(1 for r in llm_rows if (r["ai_score"] or 0) >= 70)),
        ])
    if data.llm_partial_note:
        resumen_data.append(("Nota juez LLM", data.llm_partial_note))
    for k, v in resumen_data:
        ws_resumen.append([k, _safe(v)])
    ws_resumen.column_dimensions["A"].width = 55
    ws_resumen.column_dimensions["B"].width = 15

    main_rows_sorted = sorted(main_rows, key=lambda r: (-r["score_sospecha"], r["z_tiempo_vs_grupo"] or 0))  # noqa: E501
    ws_main = wb.create_sheet("Timing y Estilo")
    headers_main = ["usuario", "problema", "score_sospecha", "un_solo_intento",
                     "intentos_antes_de_AC", "segundos_desde_su_primer_envio",
                     "z_tiempo_vs_grupo", "jplag_max_similitud", "jplag_similar_con",
                     "n_lineas", "avg_line_len", "comment_ratio", "avg_ident_len",
                     "llm_ai_score", "llm_modelo", "archivo"]
    _write_sheet(ws_main, headers_main, main_rows_sorted)

    jplag_rows_sorted = sorted(jplag_rows, key=lambda r: -r["similitud"])
    ws_jplag = wb.create_sheet("JPlag - Pares")
    headers_jplag = ["problema", "lenguaje", "usuario_a", "usuario_b", "similitud"]
    _write_sheet(ws_jplag, headers_jplag, jplag_rows_sorted)

    if data.llm_rows is not None:
        ws_llm = wb.create_sheet("LLM - Notas")
        llm_display = [
            {"usuario": r["usuario"], "problema": r["problema"], "ai_score": r["ai_score"],
             "señales": "; ".join(r["señales"]), "nota": r["nota"]}
            for r in sorted(data.llm_rows, key=lambda r: -(r["ai_score"] or -1))
        ]
        _write_sheet(ws_llm, ["usuario", "problema", "ai_score", "señales", "nota"], llm_display)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
