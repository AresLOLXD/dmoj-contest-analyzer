#!/usr/bin/env python3
"""
Analiza el export de dmoj-submission-downloader para detectar patrones de
envío sospechosos (posible uso de IA / plagio), corre JPlag, y junta todo
en un solo reporte Excel con varias hojas.

Estructura esperada del export:
nombre-concurso/
├── usuario1/
│   ├── problema_a/
│   │   ├── 1_usuario1_2024-01-15_14-30-45_AC.py
│   │   └── 2_usuario1_2024-01-15_14-35-20_WA.py
│   └── problema_b/
│       └── 1_usuario1_2024-01-15_15-10-00_AC.cpp

Requiere: pip install openpyxl --break-system-packages
(JPlag en sí requiere Java 11+ instalado por separado, solo si usas --run-jplag)

Uso básico (solo timing/estilo):
    python3 analiza_concurso.py nombre-concurso --out reporte.xlsx

Preparar Y correr JPlag, e integrar sus similitudes al Excel:
    python3 analiza_concurso.py nombre-concurso --out reporte.xlsx \\
        --jplag-out jplag_input --jplag-jar jplag.jar --run-jplag

Si ya corriste JPlag antes y solo quieres re-generar el Excel con esos
resultados existentes (sin volver a correr JPlag):
    python3 analiza_concurso.py nombre-concurso --out reporte.xlsx \\
        --jplag-out jplag_input
(sin --run-jplag: reutiliza cualquier *_resultado.jplag que ya exista ahí)
"""

import argparse
import json
import re
import shutil
import statistics
import subprocess
import zipfile
from collections import defaultdict
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    raise SystemExit(
        "Falta openpyxl. Instálalo con:\n"
        "    pip install openpyxl --break-system-packages"
    )

# ---------------------------------------------------------------------------
# Parseo común del export de dmoj-submission-downloader
# ---------------------------------------------------------------------------

FNAME_RE = re.compile(
    r"^(?P<attempt>\d+)_(?P<user>.+?)_(?P<date>\d{4}-\d{2}-\d{2})_(?P<time>\d{2}-\d{2}-\d{2})_(?P<result>[A-Z]+)\.(?P<ext>\w+)$"
)

COMMENT_PATTERNS = {
    "py": re.compile(r"^\s*#"),
    "cpp": re.compile(r"^\s*//"),
    "cc": re.compile(r"^\s*//"),
    "cxx": re.compile(r"^\s*//"),
    "java": re.compile(r"^\s*//"),
    "c": re.compile(r"^\s*//"),
}

EXT_TO_JPLAG_LANG = {
    "cpp": "cpp", "cc": "cpp", "cxx": "cpp",
    "c": "c",
    "py": "python3",
    "java": "java",
    "js": "javascript",
    "kt": "kotlin",
    "rs": "rust",
    "go": "go",
}


class Submission:
    def __init__(self, path: Path, username: str, problem: str,
                 attempt: int, dt: datetime, result: str, ext: str):
        self.path = path
        self.username = username
        self.problem = problem
        self.attempt = attempt
        self.dt = dt
        self.result = result
        self.ext = ext
        self._source = None

    @property
    def source(self) -> str:
        if self._source is None:
            self._source = self.path.read_text(errors="replace")
        return self._source

    def style_stats(self) -> dict:
        lines = [l for l in self.source.splitlines() if l.strip()]
        if not lines:
            return {"n_lines": 0, "avg_line_len": 0, "comment_ratio": 0, "avg_ident_len": 0}
        comment_re = COMMENT_PATTERNS.get(self.ext, re.compile(r"^\s*#"))
        n_comments = sum(1 for l in lines if comment_re.match(l))
        idents = re.findall(r"\b[a-zA-Z_][a-zA-Z0-9_]{1,}\b", self.source)
        keywords = {"int", "float", "double", "char", "void", "return", "include",
                    "using", "namespace", "std", "for", "while", "if", "else", "def",
                    "import", "public", "static", "class", "const", "auto", "true", "false",
                    "print", "cout", "cin", "endl", "self", "range", "len"}
        idents = [i for i in idents if i.lower() not in keywords]
        avg_ident_len = statistics.mean(len(i) for i in idents) if idents else 0
        return {
            "n_lines": len(lines),
            "avg_line_len": statistics.mean(len(l) for l in lines),
            "comment_ratio": n_comments / len(lines),
            "avg_ident_len": avg_ident_len,
        }


def parse_submissions(root: Path):
    subs = []
    for user_dir in root.iterdir():
        if not user_dir.is_dir():
            continue
        username = user_dir.name
        for prob_dir in user_dir.iterdir():
            if not prob_dir.is_dir():
                continue
            problem = prob_dir.name
            for f in prob_dir.iterdir():
                m = FNAME_RE.match(f.name)
                if not m:
                    continue
                dt = datetime.strptime(
                    f"{m.group('date')}_{m.group('time')}", "%Y-%m-%d_%H-%M-%S"
                )
                subs.append(Submission(
                    path=f,
                    username=username,
                    problem=problem,
                    attempt=int(m.group("attempt")),
                    dt=dt,
                    result=m.group("result"),
                    ext=m.group("ext").lower(),
                ))
    return subs


# ---------------------------------------------------------------------------
# Análisis de timing / estilo
# ---------------------------------------------------------------------------

def analyze_timing_style(subs):
    by_user_problem = defaultdict(list)
    for s in subs:
        by_user_problem[(s.username, s.problem)].append(s)

    first_ts_by_user = {}
    for s in subs:
        u = s.username
        if u not in first_ts_by_user or s.dt < first_ts_by_user[u]:
            first_ts_by_user[u] = s.dt

    time_to_ac_by_problem = defaultdict(list)
    rows_raw = []

    for (user, problem), items in by_user_problem.items():
        items.sort(key=lambda s: s.attempt)
        ac_subs = [s for s in items if s.result == "AC"]
        if not ac_subs:
            continue
        first_ac = ac_subs[0]
        n_attempts_before_ac = first_ac.attempt - 1
        seconds_since_user_start = (first_ac.dt - first_ts_by_user[user]).total_seconds()
        time_to_ac_by_problem[problem].append(seconds_since_user_start)

        style = first_ac.style_stats()

        rows_raw.append({
            "usuario": user,
            "problema": problem,
            "intentos_antes_de_AC": n_attempts_before_ac,
            "segundos_desde_su_primer_envio": seconds_since_user_start,
            "un_solo_intento": n_attempts_before_ac == 0,
            "n_lineas": style["n_lines"],
            "avg_line_len": round(style["avg_line_len"], 1),
            "comment_ratio": round(style["comment_ratio"], 3),
            "avg_ident_len": round(style["avg_ident_len"], 2),
            "archivo": str(first_ac.path),
            "jplag_max_similitud": None,
            "jplag_similar_con": None,
        })

    for row in rows_raw:
        times = time_to_ac_by_problem[row["problema"]]
        if len(times) > 2:
            mean_t = statistics.mean(times)
            stdev_t = statistics.pstdev(times) or 1.0
            row["z_tiempo_vs_grupo"] = round((row["segundos_desde_su_primer_envio"] - mean_t) / stdev_t, 2)
        else:
            row["z_tiempo_vs_grupo"] = None

    for row in rows_raw:
        score = 0
        if row["un_solo_intento"]:
            score += 1
        if row["z_tiempo_vs_grupo"] is not None and row["z_tiempo_vs_grupo"] < -1.0:
            score += 1
        if row["comment_ratio"] > 0.15:
            score += 1
        row["score_sospecha"] = score

    return rows_raw


# ---------------------------------------------------------------------------
# Preparación y ejecución de JPlag
# ---------------------------------------------------------------------------

def prepare_jplag_input(subs, jplag_out: Path, solo_ac: bool):
    best = {}
    for s in subs:
        lang = EXT_TO_JPLAG_LANG.get(s.ext)
        if lang is None:
            continue
        if solo_ac and s.result != "AC":
            continue
        key = (s.problem, lang, s.username)
        if key not in best or s.dt > best[key][0]:
            best[key] = (s.dt, s)

    counts = defaultdict(int)
    for (problem, lang, username), (dt, s) in best.items():
        dest_dir = jplag_out / problem / lang
        dest_dir.mkdir(parents=True, exist_ok=True)
        dest_file = dest_dir / f"{username}.{s.ext}"
        shutil.copy2(s.path, dest_file)
        counts[(problem, lang)] += 1
    return counts


def run_jplag(jplag_out: Path, counts, jplag_jar: str):
    """
    Corre JPlag con -M RUN (no abre visor, no espera Enter) para cada
    (problema, lenguaje) con 2+ usuarios. Devuelve la lista de rutas
    .jplag generadas.
    """
    result_paths = []
    for (problem, lang), n in sorted(counts.items()):
        if n < 2:
            continue
        in_dir = jplag_out / problem / lang
        result_name = jplag_out / problem / f"{lang}_resultado"
        cmd = ["java", "-jar", jplag_jar, str(in_dir), "-l", lang, "-r", str(result_name), "-M", "RUN"]
        print(f"# {problem} / {lang}  ({n} usuarios)")
        print("  " + " ".join(cmd))
        subprocess.run(cmd, check=False, stdin=subprocess.DEVNULL)
        jplag_file = result_name.with_suffix(".jplag")
        if jplag_file.exists():
            result_paths.append((problem, lang, jplag_file))
        print()
    return result_paths


def find_existing_jplag_results(jplag_out: Path):
    """Busca cualquier *_resultado.jplag ya generado bajo jplag_out."""
    results = []
    if not jplag_out.exists():
        return results
    for jplag_file in jplag_out.rglob("*_resultado.jplag"):
        problem = jplag_file.parent.name
        lang = jplag_file.stem.replace("_resultado", "")
        results.append((problem, lang, jplag_file))
    return results


# ---------------------------------------------------------------------------
# Parseo del resultado .jplag (ZIP con overview.json adentro)
#
# NOTA IMPORTANTE: el formato interno de overview.json no está documentado
# oficialmente y ha cambiado entre versiones de JPlag. Este parser usa un
# heurístico: busca recursivamente cualquier objeto JSON que tenga dos ids
# de submission + una similitud. Como armamos los archivos de entrada como
# "<usuario>.<ext>", el id de cada submission en JPlag coincide con ese
# nombre de archivo (con o sin extensión), así que no hace falta mapear
# ids a usuarios por separado.
#
# Si esto no encuentra comparaciones en tu versión de JPlag, el script
# imprime las claves de nivel superior de overview.json para que ajustemos
# el parser con datos reales, en vez de fallar en silencio.
# ---------------------------------------------------------------------------

ID_KEY_PAIRS = [
    ("firstsubmissionid", "secondsubmissionid"),
    ("firstsubmission", "secondsubmission"),
    ("submission1", "submission2"),
    ("id1", "id2"),
    ("first", "second"),
]


def _strip_ext(name: str) -> str:
    return re.sub(r"\.\w+$", "", name)


def extract_comparisons_from_json(data):
    """Devuelve lista de (id1, id2, similitud_0_a_1_o_100)."""
    found = []

    def walk(obj):
        if isinstance(obj, dict):
            lower_keys = {k.lower(): k for k in obj.keys()}
            for k1, k2 in ID_KEY_PAIRS:
                if k1 in lower_keys and k2 in lower_keys:
                    id1 = obj[lower_keys[k1]]
                    id2 = obj[lower_keys[k2]]
                    sim = None
                    if "similarities" in lower_keys and isinstance(obj[lower_keys["similarities"]], dict):
                        vals = [v for v in obj[lower_keys["similarities"]].values() if isinstance(v, (int, float))]
                        if vals:
                            sim = max(vals)
                    elif "similarity" in lower_keys and isinstance(obj[lower_keys["similarity"]], (int, float)):
                        sim = obj[lower_keys["similarity"]]
                    if isinstance(id1, str) and isinstance(id2, str) and sim is not None:
                        found.append((id1, id2, sim))
                    break
            for v in obj.values():
                walk(v)
        elif isinstance(obj, list):
            for item in obj:
                walk(item)

    walk(data)
    return found


def parse_jplag_result(problem: str, lang: str, jplag_file: Path):
    """
    Abre el .jplag (zip), lee overview.json, y devuelve una lista de dicts:
    {problema, lenguaje, usuario_a, usuario_b, similitud} con similitud en
    escala 0-100.
    """
    rows = []
    try:
        with zipfile.ZipFile(jplag_file) as zf:
            names = zf.namelist()
            overview_name = next((n for n in names if n.endswith("overview.json")), None)
            if overview_name is None:
                print(f"  [!] {jplag_file}: no se encontró overview.json dentro del zip. "
                      f"Archivos presentes: {names[:10]}{'...' if len(names) > 10 else ''}")
                return rows
            with zf.open(overview_name) as f:
                data = json.load(f)
    except Exception as e:
        print(f"  [!] No se pudo abrir/leer {jplag_file}: {e}")
        return rows

    comparisons = extract_comparisons_from_json(data)
    if not comparisons:
        print(f"  [!] No se encontraron comparaciones reconocibles en {jplag_file.name}. "
              f"Claves de nivel superior de overview.json: {list(data.keys()) if isinstance(data, dict) else type(data)}")
        return rows

    for id1, id2, sim in comparisons:
        sim_pct = sim * 100 if 0 <= sim <= 1.0 else sim
        rows.append({
            "problema": problem,
            "lenguaje": lang,
            "usuario_a": _strip_ext(id1),
            "usuario_b": _strip_ext(id2),
            "similitud": round(sim_pct, 1),
        })
    return rows


def merge_jplag_into_main(main_rows, jplag_rows):
    """Para cada (usuario, problema) del reporte principal, agrega la
    similitud máxima de JPlag y con quién, y sube el score si supera 70%."""
    best_by_user_problem = {}
    for r in jplag_rows:
        for u, other in ((r["usuario_a"], r["usuario_b"]), (r["usuario_b"], r["usuario_a"])):
            key = (u, r["problema"])
            cur = best_by_user_problem.get(key)
            if cur is None or r["similitud"] > cur[0]:
                best_by_user_problem[key] = (r["similitud"], other)

    for row in main_rows:
        key = (row["usuario"], row["problema"])
        if key in best_by_user_problem:
            sim, other = best_by_user_problem[key]
            row["jplag_max_similitud"] = sim
            row["jplag_similar_con"] = other
            if sim >= 70:
                row["score_sospecha"] += 1


# ---------------------------------------------------------------------------
# Excel multi-hoja
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
HEADER_FONT = Font(bold=True)


def _write_sheet(ws, headers, rows):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for row in rows:
        ws.append([row.get(h, "") if row.get(h, "") is not None else "" for h in headers])
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, start=1):
        max_len = max([len(str(h))] + [len(str(row.get(h, ""))) for row in rows]) if rows else len(h)
        ws.column_dimensions[get_column_letter(i)].width = min(max_len + 2, 45)


def write_excel_report(main_rows, jplag_rows, out_path: Path, n_subs, n_users, n_problems):
    wb = Workbook()

    ws_resumen = wb.active
    ws_resumen.title = "Resumen"
    ws_resumen.append(["Métrica", "Valor"])
    for cell in ws_resumen[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    n_alertas = sum(1 for r in main_rows if r["score_sospecha"] >= 2)
    n_jplag_altas = sum(1 for r in jplag_rows if r["similitud"] >= 70)
    resumen_data = [
        ("Total de submissions procesadas", n_subs),
        ("Usuarios distintos", n_users),
        ("Problemas distintos", n_problems),
        ("Casos con score_sospecha >= 2 (revisión prioritaria)", n_alertas),
        ("Pares JPlag con similitud >= 70%", n_jplag_altas),
    ]
    for k, v in resumen_data:
        ws_resumen.append([k, v])
    ws_resumen.column_dimensions["A"].width = 55
    ws_resumen.column_dimensions["B"].width = 15

    main_rows_sorted = sorted(main_rows, key=lambda r: (-r["score_sospecha"], r["z_tiempo_vs_grupo"] or 0))
    ws_main = wb.create_sheet("Timing y Estilo")
    headers_main = ["usuario", "problema", "score_sospecha", "un_solo_intento",
                     "intentos_antes_de_AC", "segundos_desde_su_primer_envio",
                     "z_tiempo_vs_grupo", "jplag_max_similitud", "jplag_similar_con",
                     "n_lineas", "avg_line_len", "comment_ratio", "avg_ident_len", "archivo"]
    _write_sheet(ws_main, headers_main, main_rows_sorted)

    jplag_rows_sorted = sorted(jplag_rows, key=lambda r: -r["similitud"])
    ws_jplag = wb.create_sheet("JPlag - Pares")
    headers_jplag = ["problema", "lenguaje", "usuario_a", "usuario_b", "similitud"]
    _write_sheet(ws_jplag, headers_jplag, jplag_rows_sorted)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("carpeta", type=Path, help="Carpeta raíz descomprimida del export del concurso")
    ap.add_argument("--out", type=Path, default=Path("reporte.xlsx"), help="Excel de salida")
    ap.add_argument("--jplag-out", type=Path, default=None,
                     help="Carpeta de trabajo de JPlag (se prepara y/o se lee de aquí)")
    ap.add_argument("--jplag-solo-ac", action="store_true",
                     help="Al preparar JPlag, usa solo el último AC de cada usuario")
    ap.add_argument("--jplag-jar", type=str, default="jplag.jar", help="Ruta al .jar de JPlag")
    ap.add_argument("--run-jplag", action="store_true",
                     help="Ejecuta JPlag (si no se pasa, se reutilizan *_resultado.jplag ya existentes en --jplag-out)")
    args = ap.parse_args()

    subs = parse_submissions(args.carpeta)
    if not subs:
        print("No se encontraron archivos que coincidan con el patrón esperado.")
        return

    n_users = len(set(s.username for s in subs))
    n_problems = len(set(s.problem for s in subs))
    print(f"{len(subs)} submissions encontradas de {n_users} usuarios en {n_problems} problemas.")

    main_rows = analyze_timing_style(subs)

    jplag_rows = []
    if args.jplag_out is not None:
        if args.run_jplag:
            counts = prepare_jplag_input(subs, args.jplag_out, args.jplag_solo_ac)
            print(f"\nEstructura de JPlag creada en: {args.jplag_out.resolve()}")
            print("\n--- JPlag ---")
            jplag_results = run_jplag(args.jplag_out, counts, args.jplag_jar)
        else:
            jplag_results = find_existing_jplag_results(args.jplag_out)
            print(f"\nReutilizando {len(jplag_results)} resultado(s) .jplag ya existentes en {args.jplag_out}")

        print("\n--- Parseando resultados de JPlag ---")
        for problem, lang, jplag_file in jplag_results:
            rows = parse_jplag_result(problem, lang, jplag_file)
            print(f"  {jplag_file.name}: {len(rows)} comparaciones extraídas")
            jplag_rows.extend(rows)

        if jplag_rows:
            merge_jplag_into_main(main_rows, jplag_rows)
        else:
            print("\n[!] No se extrajo ninguna comparación de JPlag. El Excel se genera solo con "
                  "timing/estilo. Si esperabas resultados, revisa los mensajes [!] de arriba y "
                  "compárteme las claves de overview.json impresas para ajustar el parser.")

    write_excel_report(main_rows, jplag_rows, args.out, len(subs), n_users, n_problems)
    print(f"\nReporte escrito en {args.out.resolve()}")

    top = [r for r in main_rows if r["score_sospecha"] >= 2]
    print(f"\n{len(top)} casos con score_sospecha >= 2 (revisión manual prioritaria):")
    for r in sorted(top, key=lambda r: -r["score_sospecha"])[:20]:
        print(f"  {r['usuario']:20s} {r['problema']:15s} score={r['score_sospecha']} "
              f"jplag_max_similitud={r['jplag_max_similitud']}")


if __name__ == "__main__":
    main()
