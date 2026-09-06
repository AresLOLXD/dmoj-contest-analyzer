import sys

from openpyxl import load_workbook


def read_xlsx_values(path):
    """Sheet name -> list of row dicts (header row drives keys). Order preserved."""
    wb = load_workbook(path, data_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            out[ws.title] = []
            continue
        headers = [str(h) if h is not None else "" for h in rows[0]]
        out[ws.title] = [dict(zip(headers, r, strict=False)) for r in rows[1:]]
    return out


def _run_cli(args, monkeypatch):
    monkeypatch.setattr(sys, "argv", ["dmoj-contest-analyzer", *args])
    from dmoj_contest_analyzer.cli import main
    main()


def test_cli_stdout_and_xlsx_snapshot(mini_export_tree, tmp_path, capsys, monkeypatch):
    out = tmp_path / "r.xlsx"
    _run_cli([str(mini_export_tree), "--out", str(out)], monkeypatch)
    stdout = capsys.readouterr().out

    assert "6 submissions encontradas de 3 usuarios en 2 problemas." in stdout
    assert "score_sospecha >= 2" in stdout

    values = read_xlsx_values(out)
    assert set(values) == {"Resumen", "Timing y Estilo", "JPlag - Pares"}
    ts = {(r["usuario"], r["problema"]): r for r in values["Timing y Estilo"]}
    assert ts[("userC", "p1")]["score_sospecha"] == 2
    assert ts[("userA", "p1")]["un_solo_intento"] in (True, "True", 1)
    resumen = {r["Métrica"]: r["Valor"] for r in values["Resumen"]}
    assert resumen["Usuarios distintos"] == 3
