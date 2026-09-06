import subprocess
import sys

from openpyxl import load_workbook


def test_cli_end_to_end_with_zip(flat_zip, tmp_path):
    out = tmp_path / "reporte.xlsx"
    proc = subprocess.run(
        [sys.executable, "-m", "dmoj_contest_analyzer", str(flat_zip), "--out", str(out)],
        capture_output=True, text=True,
    )
    assert proc.returncode == 0, proc.stderr
    assert out.exists()
    wb = load_workbook(out)
    assert wb.sheetnames == ["Resumen", "Timing y Estilo", "JPlag - Pares"]


def test_cli_errors_when_jplag_requested_without_jar(flat_zip, tmp_path, monkeypatch):
    monkeypatch.delenv("JPLAG_JAR", raising=False)
    out = tmp_path / "r.xlsx"
    proc = subprocess.run(
        [sys.executable, "-m", "dmoj_contest_analyzer", str(flat_zip),
         "--out", str(out), "--jplag-out", str(tmp_path / "jp"), "--run-jplag"],
        capture_output=True, text=True,
    )
    assert proc.returncode != 0
    assert "jplag" in (proc.stderr + proc.stdout).lower()
