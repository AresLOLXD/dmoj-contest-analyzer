from openpyxl import load_workbook

from dmoj_contest_analyzer.analysis import ReportData
from dmoj_contest_analyzer.report import write_excel_report
from tests.test_cli_golden import read_xlsx_values


def _main_row(**over):
    base = {
        "usuario": "userA", "problema": "p1", "score_sospecha": 2,
        "un_solo_intento": True, "intentos_antes_de_AC": 0,
        "segundos_desde_su_primer_envio": 100.0, "z_tiempo_vs_grupo": -1.5,
        "jplag_max_similitud": 80.0, "jplag_similar_con": "userC",
        "n_lineas": 10, "avg_line_len": 12.0, "comment_ratio": 0.1,
        "avg_ident_len": 4.0, "archivo": "userA/p1/x.cpp",
    }
    base.update(over)
    return base


def test_writes_three_sheets(tmp_path):
    out = tmp_path / "r.xlsx"
    write_excel_report(ReportData(
        [_main_row()],
        [{"problema": "p1", "lenguaje": "cpp", "usuario_a": "userA",
          "usuario_b": "userC", "similitud": 80.0}],
        6, 3, 2,
    ), out)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Resumen", "Timing y Estilo", "JPlag - Pares"]
    assert wb["Timing y Estilo"]["A2"].value == "userA"


def _minimal_main_row(**kw):
    row = dict(usuario="u", problema="p", score_sospecha=0, un_solo_intento=False,
               intentos_antes_de_AC=1, segundos_desde_su_primer_envio=10.0,
               z_tiempo_vs_grupo=None, jplag_max_similitud=None, jplag_similar_con=None,
               n_lineas=1, avg_line_len=1.0, comment_ratio=0.0, avg_ident_len=0.0,
               archivo="u/p/f.cpp", llm_ai_score=None, llm_modelo=None)
    row.update(kw)
    return row


def test_report_without_llm_has_empty_llm_columns(tmp_path):
    data = ReportData([_minimal_main_row()], [], 1, 1, 1)
    write_excel_report(data, tmp_path / "r.xlsx")
    v = read_xlsx_values(tmp_path / "r.xlsx")
    assert "llm_ai_score" in v["Timing y Estilo"][0]
    assert "LLM - Notas" not in v


def test_report_with_llm_adds_sheet_and_is_formula_safe(tmp_path):
    llm_rows = [{"usuario": "u", "problema": "p", "ai_score": 88,
                 "señales": ["=cmd()"], "nota": "=HYPERLINK(1)"}]
    data = ReportData([_minimal_main_row(llm_ai_score=88, llm_modelo="openai|gpt-4o")],
                      [], 1, 1, 1, llm_rows=llm_rows, llm_model="openai|gpt-4o")
    write_excel_report(data, tmp_path / "r.xlsx")
    ws = load_workbook(tmp_path / "r.xlsx")["LLM - Notas"]
    cells = [c.value for row in ws.iter_rows() for c in row]
    # formula-like strings are neutralized with a leading apostrophe and stored as text
    assert "'=HYPERLINK(1)" in cells
    assert "'=cmd()" in cells
    assert all(c.data_type != "f" for row in ws.iter_rows() for c in row)
