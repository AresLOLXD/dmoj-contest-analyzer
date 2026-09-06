import uuid

import pytest
from openpyxl import load_workbook

from dmoj_contest_analyzer.web import jobs, worker
from tests.web_conftest import make_zip

EVIL_USER = "=HYPERLINK(1)"

ZIP = {
    f"{EVIL_USER}/p1/1_x_2026-01-01_10-00-00_AC.cpp": b"int main(){}\n",
    "userB/p1/1_userB_2026-01-01_09-00-00_AC.cpp": b"int main(){}\n",
    "userC/p1/1_userC_2026-01-01_08-00-00_AC.cpp": b"int main(){}\n",
}


@pytest.mark.asyncio
async def test_attacker_dir_name_is_text_not_formula_in_xlsx(conn, settings):
    jid = uuid.uuid4().hex
    d = settings.data_dir / jid
    d.mkdir(parents=True)
    (d / "input.zip").write_bytes(make_zip(ZIP))
    jobs.create_job(conn, job_id=jid, owner="alice", model_ref=None,
                    run_jplag=False, jplag_solo_ac=False, settings=settings)

    ex = worker.make_executor(settings)
    try:
        assert await worker.process_one_job(conn, settings, ex) is True
    finally:
        ex.shutdown(wait=True)

    assert jobs.get_job(conn, jid)["status"] == "done"
    wb = load_workbook(d / "reporte.xlsx")
    ws = wb["Timing y Estilo"]
    header = [c.value for c in ws[1]]
    col = header.index("usuario") + 1
    cells = [ws.cell(row=r, column=col) for r in range(2, ws.max_row + 1)]
    evil = [c for c in cells if c.value and "HYPERLINK" in str(c.value)]
    assert evil, "expected the attacker-controlled username row"
    for c in evil:
        assert c.data_type == "s"
        assert not str(c.value).startswith("=")
